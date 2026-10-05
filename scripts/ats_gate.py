#!/usr/bin/env python3
"""ATS Gate helper for job-application-engine v2.2.0 (Phase 3C).

Runs the gate checks that can be counted, so they are the same every time:
  read         Read test (GATE-04): extract text with 2 readers, compare critical fields.
  lint         Layout lint (GATE-03): tables, text boxes, images, header/footer text, columns, fonts.
  coverage     Keyword coverage (GATE-06): must-have %, weighted %, before vs after, repeat cap.
  integrity    Locked-keyword integrity (GATE-05): every locked term present in exact form.
  hygiene      Content hygiene (GATE-09): pronouns, placeholders, date pattern consistency.
  render       Page images for the visual review (GATE-18), when pdftoppm is available.
  fingerprint  SHA-256 fingerprint of each file (GATE-14).
  all          Everything above from one config file; prints grade lines and a JSON report.
  record-pass  Freeze a PASS (GATE-14): re-runs the gate from --config (which names the job) and freezes only on PASS.
  verify-upload  Before any upload (GAP-11): exit 0 only if the file matches the job's recorded PASS.
  check-stale  After any change: mark recorded PASS files whose bytes changed as stale (edit cancels PASS).

State file: $JAE_STATE_DIR/gate_pass.json, else $CLAUDE_PROJECT_DIR/jae-state/, else ./jae-state/.

Only the Python standard library is required. Optional tools improve results:
pdftotext / pdffonts / pdftoppm (poppler), pypdf, pdfplumber.

Config file for `all` (JSON):
{
  "job": "acme-senior-pm-2026-10",                      # needed by record-pass; ties the PASS to this job
  "files": ["Alex-M-CV-Product-Manager.docx", "Alex-M-CV-Product-Manager.pdf"],
  "original_text": "path/to/original_cv.txt",          # optional, for before vs after
  "expected": {"name": "Alex M.", "email": "alex@example.com",
               "titles": ["Senior Product Manager"], "companies": ["DataFlow GmbH"],
               "dates": ["Jan 2021 – Present"], "headings": ["Summary", "Skills", "Experience", "Education"]},
  "approved_text": "path/to/approved_cv.txt",           # optional, for locked-keyword integrity
  "keywords": [{"term": "product discovery", "must": true, "weight": 3, "variants": []}],
  "date_pattern": "mon_yyyy",                            # mon_yyyy | mm_yyyy
  "allowed_fonts": ["Calibri", "Arial", "Helvetica", "Georgia"],
  "max_mb": 2, "repeat_cap": 4, "coverage_target": 0.9
}
Exit code 0 = no Must check failed (PASS or PASS WITH WAIVERS), 1 = at least one Must check failed,
2 = usage error or a missing input (no files, no keyword list, no expected name and email).
"""
import argparse, hashlib, html, json, os, re, shutil, signal, subprocess, sys, zipfile

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PRONOUNS = r"(?i)\b(i|me|my|mine|myself|we|us|our|ours|ourselves|he|him|his|himself|she|her|hers|herself)\b"
PLACEHOLDER = r"(\[[^\]]{1,40}\]|\{[^}]{1,40}\}|\bTBD\b|\bTBC\b|\bTODO\b|\bXX+\b|<[^<>\n]{2,40}>|_{3,}|(?i:lorem ipsum)|\bINSERT [A-Z]+\b)"
DATE_PATTERNS = {
    "mon_yyyy": r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]* \d{4}\b",
    "mm_yyyy": r"\b(0[1-9]|1[0-2])/\d{4}\b",
}

# ---------- text readers ----------
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

def normalize_w(xml):
    """Rewrite whatever prefix is bound to the WordprocessingML namespace to 'w:' so every
    check sees the same markup (a renamed prefix is still valid Word XML)."""
    m = re.search(r'xmlns:([A-Za-z_][\w.-]*)="' + re.escape(WNS) + '"', xml)
    if m and m.group(1) != "w":
        px = re.escape(m.group(1))
        xml = re.sub(r"(</?)" + px + r":", r"\1w:", xml)
        xml = re.sub(r"(\s)" + px + r":(?=[\w-]+=)", r"\1w:", xml)
    return xml

def main_part(z):
    """The document part named by the package relationships, not an assumed file name."""
    try:
        rels = z.read("_rels/.rels").decode("utf8", "ignore")
        for tag in re.findall(r"<Relationship\b[^>]*>", rels):
            if re.search(r'Type="[^"]*/officeDocument"', tag):
                t = re.search(r'Target="([^"]+)"', tag).group(1).lstrip("/")
                if t in z.namelist():
                    return t
    except (KeyError, AttributeError):
        pass
    return "word/document.xml"

def docx_xml(path, part=None):
    with zipfile.ZipFile(path) as z:
        part = part or main_part(z)
        return normalize_w(z.read(part).decode("utf8", "ignore")) if part in z.namelist() else ""

def docx_parts(path):
    """Every Word XML part except the main document, normalised: {name: xml}."""
    with zipfile.ZipFile(path) as z:
        main = main_part(z)
        return {n: normalize_w(z.read(n).decode("utf8", "ignore")) for n in z.namelist()
                if n.endswith(".xml") and n != main and not n.startswith(("_rels/", "docProps/", "[Content"))}

def docx_plain(path):
    """Reader A: raw text runs in document order (what simple parsers see)."""
    xml = docx_xml(path)
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab/>", "\t", xml)
    xml = re.sub(r"<w:(?:br|cr)(?:\s[^>]*)?/>", "\n", xml)  # line breaks inside a paragraph
    return html.unescape(re.sub(r"<[^>]+>", "", xml))

def docx_effective_xml(path, xml):
    """document.xml plus every style it uses (with the full basedOn chain), the default
    paragraph and character styles (found by w:default="1", not by name) and the document
    defaults where the default paragraph style doesn't override them."""
    styles = docx_xml(path, "word/styles.xml")
    blocks = re.findall(r"(<w:style\b[^>]*>)(.*?)</w:style>", styles, re.S)
    bodies, defaults = {}, set()
    for head, body in blocks:
        sid = re.search(r'w:styleId="([^"]+)"', head)
        if not sid:
            continue
        bodies[sid.group(1)] = body
        if re.search(r'w:default="(?:1|true|on)"', head):
            defaults.add(sid.group(1))
    used = set(re.findall(r'w:(?:pStyle|rStyle|tblStyle) w:val="([^"]+)"', xml)) | defaults | {"Normal"}
    seen, todo = set(), list(used)
    while todo:
        sid = todo.pop()
        if sid in seen or sid not in bodies:
            continue
        seen.add(sid)
        todo += re.findall(r'<w:basedOn w:val="([^"]+)"', bodies[sid])
    default_body = "".join(bodies.get(d, "") for d in defaults) + bodies.get("Normal", "")
    dd = re.search(r"<w:docDefaults>.*?</w:docDefaults>", styles, re.S)
    d = dd.group(0) if dd else ""
    if re.search(r'w:ascii(?:Theme)?="', default_body):
        d = re.sub(r"<w:rFonts[^>]*/>", "", d)
    if "<w:sz " in default_body:
        d = re.sub(r"<w:sz [^>]*/>", "", d)
    return xml + "".join(bodies[x] for x in seen) + d

HIDDEN = re.compile(r'<w:(?:vanish|specVanish)(?:\s+w:val="(?:true|1|on)")?\s*/>')

def near_white(xml):
    """White or near-white text colour, however the attributes are ordered (GATE-02 ban)."""
    for tag in re.findall(r"<w:color\b[^>]*/>", xml):
        if re.search(r'w:themeColor="(?:background1|bg1)"', tag):
            return True
        v = re.search(r'w:val="([0-9A-Fa-f]{6})"', tag)
        if v and all(int(v.group(1)[i:i + 2], 16) >= 0xF0 for i in (0, 2, 4)):
            return True
    return False

def docx_fonts_and_sizes(path, xml):
    """Fonts and sizes that actually apply, read from the effective formatting."""
    blob = docx_effective_xml(path, xml)
    theme = docx_xml(path, "word/theme/theme1.xml")
    theme_fonts = dict(re.findall(r'<a:(major|minor)Font>\s*<a:latin typeface="([^"]*)"', theme))
    fonts = set(re.findall(r'w:(?:ascii|hAnsi)="([^"]+)"', blob))
    for which in re.findall(r'w:(?:ascii|hAnsi)Theme="(major|minor)', blob):
        if theme_fonts.get(which):
            fonts.add(theme_fonts[which])
    sizes = [int(v) / 2 for v in re.findall(r'<w:sz(?:Cs)? w:val="(\d+)"', blob)]
    return fonts, sizes

def docx_structured(path):
    """Reader B: paragraph-level text via python-docx when available."""
    try:
        import docx  # type: ignore
        d = docx.Document(path)
        return "\n".join(p.text for p in d.paragraphs)
    except Exception:
        return None

def pdf_plain(path):
    if shutil.which("pdftotext"):
        return subprocess.run(["pdftotext", "-raw", path, "-"], capture_output=True, text=True).stdout
    try:
        import pypdf  # type: ignore
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(path).pages)
    except Exception:
        return None

def pdf_layout(path):
    if shutil.which("pdftotext"):
        return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(path) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception:
        return None

def readers(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return {"plain": docx_plain(path), "structured": docx_structured(path)}
    if ext == ".pdf":
        return {"plain": pdf_plain(path), "layout": pdf_layout(path)}
    with open(path, encoding="utf8", errors="ignore") as f:
        return {"plain": f.read()}

def norm(s):
    return re.sub(r"\s+", " ", s.replace("–", "-").replace("—", "-")).strip().lower()

# ---------- checks ----------
def check_read(path, expected):
    texts = {k: v for k, v in readers(path).items() if v}
    rows, fails = [], 0
    fields = []
    for key in ("name", "email", "phone", "linkedin"):
        if expected.get(key):
            fields.append((key, expected[key]))
    for key in ("titles", "companies", "dates", "headings"):
        for v in expected.get(key, []):
            fields.append((key, v))
    for field, value in fields:
        found = {r: norm(value) in norm(t) for r, t in texts.items()}
        ok = all(found.values()) and len(found) > 0
        fails += 0 if ok else 1
        rows.append({"field": field, "expected": value, "found": found, "match": ok})
    first10 = "\n".join((texts.get("plain") or "").splitlines()[:10])
    contact_top = all(norm(expected[k]) in norm(first10) for k in ("email",) if expected.get(k))
    return {"readers": list(texts), "reader_count": len(texts), "fields": rows,
            "failed": fails, "contact_in_first_10_lines": contact_top,
            "pass": fails == 0 and contact_top and len(texts) >= 1 and len(rows) > 0,
            "note": "" if rows else "no expected fields given: nothing was compared, so the read test cannot pass",
            "flag_single_reader": len(texts) < 2}

def check_lint(path, allowed_fonts, max_mb):
    res = {"file": os.path.basename(path), "size_mb": round(os.path.getsize(path) / 1048576, 2)}
    res["size_over_limit"] = res["size_mb"] > max_mb  # Should-level (GATE-10): reported, can be waived
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        xml = docx_xml(path)
        parts = docx_parts(path)
        hf = [x for x in parts.values() if re.search(r"<w:(?:hdr|ftr)\b", x)]
        hf_text = "".join(re.sub(r"<[^>]+>", "", x) for x in hf).strip()
        hf_images = sum(len(re.findall(r"<w:(?:drawing|pict)\b", x)) for x in hf)
        comments = any(re.search(r"<w:comments\b", x) and re.search(r"<w:t\b", x) for x in parts.values())
        with zipfile.ZipFile(path) as z:
            core = z.read("docProps/core.xml").decode("utf8", "ignore") if "docProps/core.xml" in z.namelist() else ""
        fonts, sizes = docx_fonts_and_sizes(path, xml)
        prop = lambda tag: (re.search(r"<(?:dc|cp):" + tag + r">([^<]*)<", core) or [None, ""])[1]
        pg = re.search(r'<w:pgSz [^>]*w:w="(\d+)"', xml)
        page = {"12240": "Letter", "11906": "A4", "11907": "A4"}.get(pg.group(1) if pg else "", "other")
        res.update({
            "tables": len(re.findall(r"<w:tbl\b", xml)),
            "text_boxes": len(re.findall(r"txbxContent|<w:framePr\b", xml)),
            "images": len(re.findall(r"<w:(?:drawing|pict|object)\b", xml)) + hf_images,
            "embedded_documents": len(re.findall(r"<w:altChunk\b", xml)),
            "multi_column_sections": len([c for c in re.findall(r'<w:cols\b[^>]*w:num="(\d+)"', xml) if int(c) > 1]),
            "header_footer_text": bool(hf_text),
            "comments": bool(comments),
            "tracked_changes": bool(re.search(r"<w:(?:ins|del|moveFrom|moveTo|rPrChange|pPrChange|sectPrChange|tblPrChange|trPrChange|tcPrChange)\b", xml)),
            "hidden_text": bool(HIDDEN.search(docx_effective_xml(path, xml))),
            "white_or_tiny_text": near_white(docx_effective_xml(path, xml)) or any(s < 8 for s in sizes),
            "fonts": sorted(fonts),
            "fonts_outside_allowed": sorted(f for f in fonts if allowed_fonts and f not in allowed_fonts),
            "body_sizes_outside_10_12": sorted({s for s in sizes if s < 10 or (12 < s < 14)}),
            "page_size": page,
            "properties": {"title": prop("title"), "author": prop("creator"), "last_modified_by": prop("lastModifiedBy")},
        })
        bad_author = (not res["properties"]["author"]) or "python-docx" in (res["properties"]["author"] + prop("description")).lower()
        res["properties_need_fixing"] = bad_author or not res["properties"]["title"]
        must_zero = ["tables", "text_boxes", "images", "embedded_documents", "multi_column_sections"]
        res["pass"] = (all(res[k] == 0 for k in must_zero) and not res["header_footer_text"]
                       and not res["comments"] and not res["tracked_changes"] and not res["hidden_text"]
                       and not res["white_or_tiny_text"] and not res["fonts_outside_allowed"]
                       and not res["body_sizes_outside_10_12"]
                       and not res["properties_need_fixing"])
    elif ext == ".pdf":
        res["text_layer"] = bool((pdf_plain(path) or "").strip())
        res["checks_not_run"] = []
        if shutil.which("pdffonts"):
            out = subprocess.run(["pdffonts", path], capture_output=True, text=True).stdout.splitlines()[2:]
            res["fonts_not_embedded"] = [l.split()[0] for l in out if len(l.split()) > 4 and l.split()[-5] == "no"]
        else:
            res["checks_not_run"].append("embedded fonts (pdffonts missing)")
        if shutil.which("pdfimages"):
            out = subprocess.run(["pdfimages", "-list", path], capture_output=True, text=True).stdout.splitlines()[2:]
            res["images"] = len(out)
        res["invisible_text"] = None
        try:
            import pypdf  # type: ignore
            reader = pypdf.PdfReader(path)
            def streams(pg):
                c = pg.get("/Contents")
                c = c.get_object() if c is not None else None
                items = c if isinstance(c, list) else ([c] if c is not None else [])
                return b"".join(x.get_object().get_data() for x in items)
            data = b"".join(streams(pg) for pg in reader.pages)
            res["invisible_text"] = bool(re.search(rb"(?<![\d.])3\s+Tr\b", data))
            if "images" not in res:
                res["images"] = sum(len(getattr(pg, "images", [])) for pg in reader.pages)
        except Exception:
            res["checks_not_run"].append("invisible text (pypdf missing or unreadable PDF)")
        if "images" not in res:
            res["checks_not_run"].append("images (pdfimages and pypdf missing)")
        res["pass"] = (res["text_layer"] and not res.get("fonts_not_embedded") and res.get("images", 0) == 0
                       and not res["invisible_text"])
    else:
        res["pass"] = False
        res["error"] = "Only .docx and .pdf are submission formats."
    return res

def count_term(text, term, variants=()):
    t = norm(text)
    return sum(len(re.findall(r"(?<!\w)" + re.escape(norm(v)) + r"(?!\w)", t)) for v in [term, *variants])

def check_coverage(text, keywords, original=None, cap=4, target=0.9):
    def score(tx):
        must = [k for k in keywords if k.get("must")]
        hit_must = [k for k in must if count_term(tx, k["term"], k.get("variants", [])) > 0]
        tot_w = sum(k.get("weight", 1) for k in keywords) or 1
        hit_w = sum(k.get("weight", 1) for k in keywords if count_term(tx, k["term"], k.get("variants", [])) > 0)
        return (len(hit_must) / len(must) if must else (1.0 if keywords else 0.0)), hit_w / tot_w
    must_after, w_after = score(text)
    rows = [{"term": k["term"], "must": bool(k.get("must")), "count": count_term(text, k["term"], k.get("variants", []))} for k in keywords]
    over = [r for r in rows if r["count"] > cap]
    res = {"must_have_coverage": round(must_after, 3), "weighted_coverage": round(w_after, 3),
           "over_cap": over, "missing_must": [r["term"] for r in rows if r["must"] and r["count"] == 0],
           "rows": rows, "pass": must_after >= target and not over}
    if original is not None:
        mb, wb = score(original)
        res.update({"must_have_coverage_before": round(mb, 3), "weighted_coverage_before": round(wb, 3)})
    return res

def check_integrity(text, keywords, approved=None):
    """Exact-form check on the extracted text (catches special dashes, split ligatures, smart quotes).
    With the approved CV text: every locked term found there must appear unchanged in the file text.
    Without it: every term that appears in loose form must also appear in exact form."""
    missing = []
    for k in keywords:
        if not k.get("locked", True):
            continue
        term = k["term"]
        expected_present = (term.lower() in approved.lower()) if approved is not None else count_term(text, term) > 0
        if expected_present and term.lower() not in text.lower():
            missing.append(term)
    return {"missing_exact_form": missing, "checked_against": "approved text" if approved is not None else "loose matches", "pass": not missing}

def check_headings(text, headings):
    """REQ-24: each section heading appears once, as its own line, in the agreed order."""
    if not headings:
        return {"duplicate_headings": [], "out_of_order": False, "checked": False}
    lines = [norm(l) for l in text.splitlines() if l.strip()]
    counts = {h: lines.count(norm(h)) for h in headings}
    firsts = [lines.index(norm(h)) for h in headings if norm(h) in lines]
    return {"duplicate_headings": sorted(h for h, c in counts.items() if c > 1),
            "out_of_order": firsts != sorted(firsts), "checked": True}

NUMERAL_I = re.compile(r"(Phase|Tier|Level|Stage|Series|Grade|Type|Class|Part|Year|Act|Step|Gen|War)\s+$", re.I)

def not_a_pronoun(text, m):
    """'I' inside I/O, I-9 or 'Phase I' is not a first-person pronoun; neither is the country 'US'."""
    if m.group(0) == "US":
        return True
    if m.group(0) not in ("I", "i"):
        return False
    before, after = text[max(0, m.start() - 12):m.start()], text[m.end():m.end() + 1]
    return after in ("/", "-") or before.endswith(("/", "-")) or bool(NUMERAL_I.search(before))

def check_hygiene(text, date_pattern="mon_yyyy", headings=None):
    pron = sorted(set(m.group(0) for m in re.finditer(PRONOUNS, text) if not not_a_pronoun(text, m)))
    ph = sorted(set(m.group(0) for m in re.finditer(PLACEHOLDER, text)))
    other = [p for k, p in DATE_PATTERNS.items() if k != date_pattern]
    mixed = any(re.search(p, text) for p in other) and re.search(DATE_PATTERNS[date_pattern], text)
    hd = check_headings(text, headings or [])
    return {"pronouns": pron, "placeholders": ph, "mixed_date_formats": bool(mixed), **hd,
            "pass": not pron and not ph and not mixed and not hd["duplicate_headings"] and not hd["out_of_order"]}

def render(path, outdir):
    if not path.lower().endswith(".pdf") or not shutil.which("pdftoppm"):
        return {"rendered": [], "note": "Render needs a PDF and pdftoppm. Convert DOCX to PDF first (e.g. soffice --headless --convert-to pdf)."}
    os.makedirs(outdir, exist_ok=True)
    stem = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0])
    subprocess.run(["pdftoppm", "-png", "-r", "110", path, stem], check=False)
    return {"rendered": sorted(os.path.join(outdir, f) for f in os.listdir(outdir) if f.endswith(".png"))}

def fingerprint(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()

# ---------- frozen PASS state (GATE-14, GAP-11) ----------
def state_path():
    base = os.environ.get("JAE_STATE_DIR") or os.path.join(os.environ.get("CLAUDE_PROJECT_DIR") or os.getcwd(), "jae-state")
    return os.path.join(base, "gate_pass.json")

def load_state():
    p = state_path()
    if os.path.exists(p):
        with open(p, encoding="utf8") as f:
            return json.load(f)
    return {"jobs": {}}

def save_state(state):
    p = state_path()
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf8") as f:
        json.dump(state, f, indent=2)

def resolve_paths(cfg, config_path):
    """Relative paths in a gate config: as given if they exist from here, else next to the config."""
    base = os.path.dirname(os.path.abspath(config_path))
    fix = lambda f: f if (os.path.isabs(f) or os.path.exists(f)) else os.path.join(base, f)
    for k in ("approved_text", "original_text"):
        if cfg.get(k):
            cfg[k] = os.path.abspath(fix(cfg[k]))
    cfg["files"] = [os.path.abspath(fix(f)) for f in cfg.get("files", [])]
    if isinstance(cfg.get("keywords"), str):  # a path to a keywords file, as `coverage` and `integrity` accept
        kp = fix(cfg["keywords"])
        try:
            cfg["keywords"] = json.load(open(kp, encoding="utf8"))
        except (OSError, ValueError):
            print(f"ATS Gate cannot start (GATE-01 entry criteria): keywords file not readable: {kp}", file=sys.stderr)
            sys.exit(2)
    cfg["_config_path"] = os.path.abspath(config_path)
    return cfg

def keywords_changed(entry):
    """True when the job's keyword list changed after its PASS (the PASS no longer proves anything)."""
    cp = entry.get("config_path")
    if not cp or not os.path.exists(cp):
        return False
    try:
        return keywords_fingerprint(json.load(open(cp, encoding="utf8")).get("keywords", [])) != entry.get("keywords_sha256")
    except Exception:
        return True

def keywords_fingerprint(kws):
    return hashlib.sha256(json.dumps(kws, sort_keys=True).encode("utf8")).hexdigest()

def record_pass(job, cfg, waived=False):
    """Freeze a PASS (GATE-14). Re-runs every counted check from the gate config itself, so a
    PASS can't come from a hand-written or stale report, and ties it to this job's keyword list."""
    if cfg.get("job") != job:
        return {"recorded": False, "reason": f"the gate config is for job {cfg.get('job')!r}, not {job!r}; add \"job\": \"{job}\" to the config and re-run"}
    rep = run_all(cfg)
    ok = rep["verdict"] == "PASS" or (waived and rep["verdict"] == "PASS WITH WAIVERS")
    if not ok:
        hint = " (add --waived only after the person accepts each waiver)" if rep["verdict"] == "PASS WITH WAIVERS" else ""
        return {"recorded": False, "reason": f"gate verdict is {rep['verdict']} ({rep['grade_line']}), not PASS{hint}"}
    state = load_state()
    state["jobs"][job] = {"files": {os.path.abspath(f): fingerprint(f) for f in cfg["files"]}, "stale": [],
                          "keywords_sha256": keywords_fingerprint(cfg.get("keywords", [])),
                          "verdict": rep["verdict"], "grade_line": rep["grade_line"],
                          "config_path": cfg.get("_config_path"),
                          "recorded_at": __import__("datetime").datetime.now().isoformat(timespec="seconds")}
    state["active_job"] = job  # the job being applied to now; uploads without --job are checked against it
    save_state(state)
    return {"recorded": True, "job": job, "state_file": state_path(), "files": list(state["jobs"][job]["files"])}

def verify_upload(path, job=None):
    state, fp, ap = load_state(), fingerprint(path), os.path.abspath(path)
    job = job or state.get("active_job")  # without --job, the most recent PASS's job is the one in progress
    jobs = {job: state["jobs"].get(job)} if job else state["jobs"]
    for j, entry in jobs.items():
        if not entry:
            continue
        for f, h in entry["files"].items():
            if h == fp:
                if f in entry.get("stale", []):
                    return {"pass": False, "reason": f"{os.path.basename(path)} matches a PASS for job {j}, but that PASS was cancelled by a later edit. Re-run the ATS Gate."}
                if keywords_changed(entry):
                    return {"pass": False, "reason": f"the keyword list for job {j} changed after this PASS. Re-run the ATS Gate."}
                return {"pass": True, "job": j, "file": os.path.basename(path)}
    why = f"no ATS Gate PASS for job {job}" if job and not state["jobs"].get(job) else "this file is not the file that passed the ATS Gate (fingerprint differs)"
    return {"pass": False, "reason": f"{os.path.basename(path)}: {why}. Only the tested file may be uploaded (GAP-11)."}

def check_stale():
    state, newly = load_state(), []
    for j, entry in state["jobs"].items():
        kw_changed = keywords_changed(entry)
        for f, h in entry["files"].items():
            if f in entry.get("stale", []):
                continue
            if kw_changed or not os.path.exists(f) or fingerprint(f) != h:
                entry.setdefault("stale", []).append(f)
                newly.append({"job": j, "file": os.path.basename(f)})
    if newly:
        save_state(state)
    return {"newly_stale": newly, "pass": True}

# ---------- grade line ----------
def grade(must_fails, should_fails, warnings):
    if must_fails:
        return 1 if must_fails > 1 else 2
    if warnings >= 3:
        return 3
    return 4 if should_fails or warnings else 5

def verdict_for(must_fails, g):
    """One verdict that always agrees with the grade: FAIL on any Must failure; PASS WITH WAIVERS
    when only warnings remain (the person must accept each one); PASS only at 5/5."""
    return "FAIL" if must_fails else ("PASS WITH WAIVERS" if g < 5 else "PASS")

def validate_config(cfg):
    """Fail fast with a plain message instead of passing on missing inputs (GATE-01)."""
    problems = []
    files = cfg.get("files") or []
    if not files:
        problems.append("config has no 'files' to test")
    for f in files:
        if not os.path.exists(f):
            problems.append(f"file not found: {f}")
        elif os.path.splitext(f)[1].lower() not in (".docx", ".pdf"):
            problems.append(f"not a submission format (.docx or .pdf): {f}")
    if cfg.get("keywords") and not (isinstance(cfg["keywords"], list) and all(isinstance(k, dict) and k.get("term") for k in cfg["keywords"])):
        problems.append("'keywords' must be a list of {\"term\": ..., \"must\": true/false} entries, or a path to such a file")
    elif not cfg.get("keywords"):
        problems.append("config has no 'keywords' list (from the Step 2A job analysis)")
    exp = cfg.get("expected") or {}
    if not (exp.get("name") and exp.get("email")):
        problems.append("config 'expected' needs at least the person's name and email for the read test")
    if problems:
        print("ATS Gate cannot start (GATE-01 entry criteria): " + "; ".join(problems), file=sys.stderr)
        sys.exit(2)

def run_all(cfg):
    validate_config(cfg)
    report, must_fail, should_fail, warn, reasons = {"files": {}}, 0, 0, 0, []
    kws = cfg.get("keywords", [])
    approved = None
    if cfg.get("approved_text") and os.path.exists(cfg["approved_text"]):
        approved = open(cfg["approved_text"], encoding="utf8", errors="ignore").read()
    original = None
    if cfg.get("original_text") and os.path.exists(cfg["original_text"]):
        original = open(cfg["original_text"], encoding="utf8", errors="ignore").read()
    for path in cfg["files"]:
        r = {"fingerprint": fingerprint(path)}
        r["lint"] = check_lint(path, cfg.get("allowed_fonts", []), cfg.get("max_mb", 2))
        r["read"] = check_read(path, cfg.get("expected", {}))
        text = readers(path).get("plain") or ""
        r["integrity"] = check_integrity(text, kws, approved)
        r["coverage"] = check_coverage(text, kws, original, cfg.get("repeat_cap", 4), cfg.get("coverage_target", 0.9))
        r["hygiene"] = check_hygiene(text, cfg.get("date_pattern", "mon_yyyy"), cfg.get("expected", {}).get("headings"))
        name = os.path.basename(path)
        for check in ("lint", "read", "integrity", "coverage", "hygiene"):
            if not r[check]["pass"]:
                must_fail += 1
                reasons.append(f"{name}: {check} failed")
        if not r["read"]["fields"]:
            warn += 1
            reasons.append(f"{name}: no expected fields given, read test compared nothing")
        if r["lint"].get("size_over_limit"):
            should_fail += 1
            reasons.append(f"{name}: file is {r['lint']['size_mb']} MB (limit {cfg.get('max_mb', 2)} MB)")
        if r["integrity"].get("checked_against") == "loose matches":
            warn += 1
            reasons.append(f"{name}: no approved CV text given, locked keywords checked against loose matches only")
        if r["read"].get("flag_single_reader"):
            warn += 1
            reasons.append(f"{name}: only 1 text reader available")
        report["files"][name] = r
    for f in cfg["files"]:
        miss = report["files"][os.path.basename(f)]["lint"].get("checks_not_run")
        if miss:
            warn += 1
            reasons.append(f"{os.path.basename(f)}: not checked — {', '.join(miss)}")
    # A DOCX and its PDF must say the same thing: a PDF with extra or hidden words fails (GATE-02).
    docx_by_stem = {os.path.splitext(os.path.basename(f))[0]: f for f in cfg["files"] if f.lower().endswith(".docx")}
    for f in cfg["files"]:
        stem = os.path.splitext(os.path.basename(f))[0]
        if f.lower().endswith(".pdf") and stem in docx_by_stem:
            a = set(re.findall(r"\w+", norm(docx_plain(docx_by_stem[stem]))))
            b = set(re.findall(r"\w+", norm(pdf_plain(f) or "")))
            sim = len(a & b) / max(len(a | b), 1)
            report["files"][os.path.basename(f)]["docx_pdf_word_match"] = round(sim, 3)
            if sim < 0.9:
                must_fail += 1
                reasons.append(f"{os.path.basename(f)}: text differs from {os.path.basename(docx_by_stem[stem])} ({round(sim * 100)}% of words match, needs 90%)")
    stems = {os.path.splitext(os.path.basename(f))[0] for f in cfg["files"] if f.lower().endswith(".docx")}
    for f in cfg["files"]:
        if f.lower().endswith(".pdf") and os.path.splitext(os.path.basename(f))[0] not in stems:
            warn += 1  # PDF text can't show tables or header text; the source DOCX must be gated with it
            reasons.append(f"{os.path.basename(f)}: tested without its source DOCX, so tables and header text can't be seen")
    g = grade(must_fail, should_fail, warn)
    verdict = verdict_for(must_fail, g)
    reason = "; ".join(reasons) if reasons else "all counted checks passed"
    stop = "Stopping." if g <= 3 else "Not stopping."
    report["grade"], report["verdict"] = g, verdict
    report["grade_line"] = f"ATS Gate (counted checks): {g}/5 · {reason}. {stop}"
    report["note"] = "Visual review (GATE-18), evidence trace (GATE-07), knockout alignment (GATE-08), platform and country profiles (GATE-10, GATE-11) are judged by the engine, not this script."
    return report

def main():
    if hasattr(signal, "SIGPIPE"):
        signal.signal(signal.SIGPIPE, signal.SIG_DFL)  # quiet exit when output is piped to head
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for c in ("read", "lint", "hygiene", "fingerprint"):
        p = sub.add_parser(c); p.add_argument("file")
        if c == "read": p.add_argument("--expected", help="JSON file with expected fields")
        if c == "lint": p.add_argument("--max-mb", type=float, default=2); p.add_argument("--fonts", default="Calibri,Arial,Helvetica,Georgia")
        if c == "hygiene":
            p.add_argument("--date-pattern", default="mon_yyyy", choices=list(DATE_PATTERNS))
            p.add_argument("--headings", default="", help="comma-separated section headings in their agreed order")
    p = sub.add_parser("coverage"); p.add_argument("file"); p.add_argument("--keywords", required=True); p.add_argument("--original"); p.add_argument("--cap", type=int, default=4); p.add_argument("--target", type=float, default=0.9)
    p = sub.add_parser("integrity"); p.add_argument("file"); p.add_argument("--keywords", required=True); p.add_argument("--approved", help="approved CV text file")
    p = sub.add_parser("render"); p.add_argument("file"); p.add_argument("--out", default="gate_pages")
    p = sub.add_parser("all"); p.add_argument("config"); p.add_argument("--out", help="also write the report as pure JSON to this file (for the evidence pack)")
    p = sub.add_parser("record-pass"); p.add_argument("--job", required=True); p.add_argument("--config", required=True, help="the gate config for this job (must contain \"job\"); the gate is re-run from it"); p.add_argument("--waived", action="store_true", help="also accept PASS WITH WAIVERS, after the person accepted each waiver")
    p = sub.add_parser("verify-upload"); p.add_argument("file"); p.add_argument("--job")
    sub.add_parser("check-stale")
    a = ap.parse_args()
    load = lambda f: json.load(open(f, encoding="utf8"))
    text_of = lambda f: readers(f).get("plain") or ""
    if a.cmd == "read": out = check_read(a.file, load(a.expected) if a.expected else {})
    elif a.cmd == "lint": out = check_lint(a.file, [x for x in a.fonts.split(",") if x], a.max_mb)
    elif a.cmd == "hygiene": out = check_hygiene(text_of(a.file), a.date_pattern, [h.strip() for h in a.headings.split(",") if h.strip()])
    elif a.cmd == "fingerprint": out = {"file": a.file, "sha256": fingerprint(a.file)}
    elif a.cmd == "coverage":
        orig = text_of(a.original) if a.original else None
        out = check_coverage(text_of(a.file), load(a.keywords), orig, a.cap, a.target)
    elif a.cmd == "integrity": out = check_integrity(text_of(a.file), load(a.keywords), open(a.approved, encoding="utf8").read() if a.approved else None)
    elif a.cmd == "render": out = render(a.file, a.out)
    elif a.cmd == "record-pass":
        cfg = resolve_paths(load(a.config), a.config)
        out = record_pass(a.job, cfg, a.waived); out["pass"] = out["recorded"]
    elif a.cmd == "verify-upload": out = verify_upload(a.file, a.job)
    elif a.cmd == "check-stale": out = check_stale()
    else:
        out = run_all(resolve_paths(load(a.config), a.config))
        if a.out:
            with open(a.out, "w", encoding="utf8") as fh:
                json.dump(out, fh, indent=2, ensure_ascii=False)
        print(out["grade_line"])
    print(json.dumps(out, indent=2, ensure_ascii=False))
    if "pass" not in out and "verdict" not in out:
        sys.exit(0)  # informational commands (fingerprint, render) succeed when they print a result
    sys.exit(0 if out.get("pass", out.get("verdict") in ("PASS", "PASS WITH WAIVERS")) else 1)

if __name__ == "__main__":
    main()
